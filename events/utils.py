import logging
import openpyxl
from django.utils import timezone
import datetime

logger = logging.getLogger(__name__)


def write_xls(sheet_name, head, data_list):
    """write listed data into excel
    """

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
    except Exception as e:
        logger.error(e)
        return None

    ws.title = sheet_name

    row_num = 0

    # write table head
    for col_num in range(len(head)):
        c = ws.cell(row=row_num + 1, column=col_num + 1)
        c.value = head[col_num]

    # write table data
    for row in data_list:
        row_num += 1
        for col_num in range(len(row)):
            c = ws.cell(row=row_num + 1, column=col_num + 1)
            c.value = row[col_num]

    return wb


def utc_to_local(dt):
    # change from UTC timezone to current seahub timezone
    tz = timezone.get_default_timezone()
    utc = dt.replace(tzinfo=datetime.timezone.utc)
    local = timezone.make_naive(utc, tz)

    return local

def generate_file_audit_event_type(e):
    event_type_dict = {
        'file-download-web': ('web', ''),
        'file-download-share-link': ('share-link', ''),
        'file-download-api': ('API', e.device),
        'repo-download-sync': ('download-sync', e.device),
        'repo-upload-sync': ('upload-sync', e.device),
        'seadrive-download-file': ('seadrive-download', e.device),
    }

    if e.etype not in event_type_dict:
        event_type_dict[e.etype] = (e.etype, e.device if e.device else '')

    return event_type_dict[e.etype]
