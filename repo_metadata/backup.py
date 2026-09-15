import json
from datetime import datetime, timezone

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


FORMAT_VERSION = 'metadata-backup-v1'
MAIN_TABLE_ID = '0001'
MAIN_TABLE_NAME = 'Table1'
TAGS_TABLE_NAME = 'tags'
PAGE_SIZE = 1000

SYSTEM_COLUMN_KEYS = {'_creator', '_ctime', '_last_modifier', '_mtime', '_dtime'}
PRESERVED_COLUMN_KEYS = {
    '_file_creator', '_file_ctime', '_file_modifier', '_file_mtime',
    '_parent_dir', '_name', '_is_dir', '_file_type', '_obj_id', '_size',
    '_suffix', '_file_details',
}
CONDITIONAL_COLUMN_KEYS = {
    '_ai_summary', '_ai_summary_mtime', '_location', '_location_translated',
    '_keywords', '_capture_time',
}
IGNORED_COLUMN_KEYS = {
    '_face_vectors', '_face_links', '_excluded_face_links', '_included_face_links',
}
JSON_TYPES = {'multiple-select', 'collaborator', 'geolocation', 'digital-sign'}
SUPPORTED_COLUMN_TYPES = {
    'checkbox', 'number', 'rate', 'duration', 'text', 'long-text', 'url', 'email',
    'date', 'collaborator', 'geolocation', 'digital-sign', 'single-select', 'multiple-select',
}

CELL_CHUNK_SIZE = 30000
SCHEMA_HEADERS = ['table_id', 'table_name', 'column_index', 'key', 'name', 'type', 'mode', 'data_part', 'data']
VALUE_HEADERS = ['table_id', 'row_id', 'column_key', 'part', 'value']
LINK_DEFINITION_HEADERS = [
    'link_id', 't1_id', 't2_id',
    't1_column_key', 't1_column_name', 't1_display_column_key',
    't2_column_key', 't2_column_name', 't2_display_column_key',
]
LINK_RELATION_HEADERS = ['row_id', 'other_row_id']
LINK_HEADERS = LINK_DEFINITION_HEADERS + LINK_RELATION_HEADERS


class MetadataBackupError(Exception):
    pass


def export_metadata_backup(metadata_server_api, repo_id, repo_name, settings, target_path):
    metadata = metadata_server_api.get_metadata(repo_id)
    tables = metadata.get('tables') or []
    main = next((table for table in tables if table.get('id') == MAIN_TABLE_ID), None)
    if not main:
        raise MetadataBackupError('Metadata table not found')
    tags = next((table for table in tables if table.get('name') == TAGS_TABLE_NAME), None)

    included_tables = [main]
    if tags:
        included_tables.append(tags)
    schemas = {table['id']: _exported_columns(table) for table in included_tables}
    table_rows = {
        table['id']: _query_all_rows(metadata_server_api, repo_id, table, schemas[table['id']])
        for table in included_tables
    }

    workbook = openpyxl.Workbook(write_only=True)
    schema_sheet = workbook.create_sheet('_Schema')
    links_sheet = workbook.create_sheet('_Links')
    views_sheet = workbook.create_sheet('_Views')
    manifest_sheet = workbook.create_sheet('_Manifest')
    values_sheet = workbook.create_sheet('_Values')
    schema_sheet.sheet_state = 'hidden'
    links_sheet.sheet_state = 'hidden'
    views_sheet.sheet_state = 'hidden'
    manifest_sheet.sheet_state = 'hidden'
    values_sheet.sheet_state = 'hidden'
    _write_header(values_sheet, VALUE_HEADERS)

    for table in included_tables:
        sheet_name = 'Metadata' if table['id'] == MAIN_TABLE_ID else 'Tags'
        sheet = workbook.create_sheet(sheet_name, 0 if table['id'] == MAIN_TABLE_ID else 1)
        columns = schemas[table['id']]
        _write_data_sheet(sheet, table['id'], columns, table_rows[table['id']], values_sheet)

    _write_header(schema_sheet, SCHEMA_HEADERS)
    for table in included_tables:
        for index, column in enumerate(schemas[table['id']], 1):
            data = json.dumps(column.get('data'), ensure_ascii=False, separators=(',', ':'))
            for part, chunk in enumerate(_chunks(data), 1):
                schema_sheet.append([
                    table['id'], table['name'], index, column['key'], column['name'],
                    column['type'], _column_mode(table, column), part,
                    _string_cell(schema_sheet, chunk),
                ])

    _write_links(links_sheet, included_tables, table_rows)
    _write_key_values(views_sheet, settings)
    _write_key_values(manifest_sheet, {
        'format': FORMAT_VERSION,
        'repo_id': repo_id,
        'repo_name': repo_name,
        'created_at': datetime.now(timezone.utc).isoformat(),
        'metadata_rows': len(table_rows[main['id']]),
        'tag_rows': len(table_rows.get(tags['id'], [])) if tags else 0,
    })
    workbook.save(target_path)


def import_metadata_backup(source_path, repo_id):
    workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=False)
    try:
        required = {'Metadata', '_Schema', '_Links', '_Views', '_Manifest', '_Values'}
        if not required.issubset(workbook.sheetnames):
            raise MetadataBackupError('Invalid metadata backup workbook')

        manifest = _read_key_values(workbook['_Manifest'])
        if manifest.get('format') != FORMAT_VERSION:
            raise MetadataBackupError('Unsupported metadata backup format')
        if manifest.get('repo_id') != repo_id:
            raise MetadataBackupError('The backup belongs to another library')

        schema = _read_schema(workbook['_Schema'])
        main_columns = schema.get(MAIN_TABLE_ID)
        if not main_columns:
            raise MetadataBackupError('Metadata schema not found')
        _validate_schema(schema)
        long_values = _read_long_values(workbook['_Values'])
        main_rows = _read_data_sheet(workbook['Metadata'], main_columns, long_values)
        _validate_unique_row_ids(main_rows, MAIN_TABLE_NAME)

        tags_schema = next((columns for table_id, columns in schema.items()
                            if columns[0]['table_name'] == TAGS_TABLE_NAME), None)
        tables = []
        tag_rows = []
        if tags_schema:
            if 'Tags' not in workbook.sheetnames:
                raise MetadataBackupError('Tags sheet not found')
            tag_rows = _read_data_sheet(workbook['Tags'], tags_schema, long_values)
            _validate_unique_row_ids(tag_rows, TAGS_TABLE_NAME)
            tables.append({
                'id': tags_schema[0]['table_id'],
                'name': TAGS_TABLE_NAME,
                'columns': [_api_column(column) for column in tags_schema if column['key'] != '_id'],
                'rows': tag_rows,
            })

        restore_columns = [
            column for column in main_columns
            if column['mode'] in {'restore', 'conditional'} and column['key'] != '_id'
        ]
        restore_keys = {column['key'] for column in restore_columns}
        payload_rows = []
        for row in main_rows:
            payload_row = {'_id': row['_id'], '_obj_id': row.get('_obj_id')}
            payload_row.update({key: value for key, value in row.items() if key in restore_keys})
            payload_rows.append(payload_row)

        settings = _read_key_values(workbook['_Views'])
        return {
            'payload': {
                'table_id': MAIN_TABLE_ID,
                'preserved_column_keys': [
                    column['key'] for column in main_columns if column['mode'] == 'preserve'
                ],
                'match_column_key': '_obj_id',
                'conditional_column_keys': [
                    column['key'] for column in restore_columns if column['mode'] == 'conditional'
                ],
                'columns': [_api_column(column) for column in restore_columns],
                'rows': payload_rows,
                'tables': tables,
                'links': _read_links(workbook['_Links']),
            },
            'settings': settings,
            'summary': {
                'metadata_rows': len(main_rows),
                'columns': len(restore_columns),
                'tags': len(tag_rows),
            },
        }
    finally:
        workbook.close()


def _query_all_rows(metadata_server_api, repo_id, table, exported_columns):
    rows = []
    escaped_name = table['name'].replace('`', '``')
    columns = list(exported_columns)
    columns.extend(
        column for column in table.get('columns') or []
        if column.get('type') == 'link' and column.get('key') not in IGNORED_COLUMN_KEYS
    )
    fields = ', '.join(f'`{column["name"].replace("`", "``")}`' for column in columns)
    while True:
        result = metadata_server_api.query_rows(
            repo_id,
            f'SELECT {fields} FROM `{escaped_name}` LIMIT {len(rows)}, {PAGE_SIZE}'
        ).get('results') or []
        rows.extend(result)
        if len(result) < PAGE_SIZE:
            return rows


def _exported_columns(table):
    columns = []
    for column in table.get('columns') or []:
        key = column.get('key')
        if column.get('type') == 'link' or key in SYSTEM_COLUMN_KEYS or key in IGNORED_COLUMN_KEYS:
            continue
        columns.append(column)
    if not columns or columns[0].get('key') != '_id':
        raise MetadataBackupError(f'Invalid schema for table {table.get("name")}')
    return columns


def _column_mode(table, column):
    if column['key'] == '_id':
        return 'identity'
    if table['id'] != MAIN_TABLE_ID:
        return 'restore'
    if column['key'] in PRESERVED_COLUMN_KEYS:
        return 'preserve'
    if column['key'] in CONDITIONAL_COLUMN_KEYS:
        return 'conditional'
    return 'restore'


def _write_data_sheet(sheet, table_id, columns, rows, values_sheet):
    _write_header(sheet, [column['name'] for column in columns])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{len(rows) + 1}'
    sheet.column_dimensions['A'].hidden = True
    for index, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 24 if column['key'] != '_id' else 4
    for row in rows:
        values = []
        for column in columns:
            value = _encode_value(column, row.get(column['name']))
            if isinstance(value, str) and len(value) > CELL_CHUNK_SIZE:
                for part, chunk in enumerate(_chunks(value), 1):
                    values_sheet.append([
                        table_id, row['_id'], column['key'], part,
                        _string_cell(values_sheet, chunk),
                    ])
                value = value[:CELL_CHUNK_SIZE]
            values.append(_string_cell(sheet, value))
        sheet.append(values)


def _write_header(sheet, values):
    cells = []
    for value in values:
        cell = _string_cell(sheet, value)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='4F81BD')
        cells.append(cell)
    sheet.append(cells)


def _string_cell(sheet, value):
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.data_type = 's'
    return cell


def _encode_value(column, value):
    if value is None:
        return None
    if column['type'] in JSON_TYPES:
        return json.dumps(value, ensure_ascii=False, separators=(',', ':'))
    return value


def _write_links(sheet, tables, table_rows):
    _write_header(sheet, LINK_HEADERS)
    table_by_id = {table['id']: table for table in tables}
    link_columns = {}
    for table in tables:
        for column in table.get('columns') or []:
            if column.get('type') != 'link':
                continue
            data = column.get('data') or {}
            link_id = data.get('link_id')
            if link_id:
                link_columns.setdefault(link_id, []).append((table, column))

    for link_id, columns in link_columns.items():
        forward = next((item for item in columns
                        if (item[1].get('data') or {}).get('table_id') == item[0]['id']
                        and not (item[1].get('data') or {}).get('is_linked_back')), None)
        backward = next((item for item in columns if item != forward), None)
        if not forward or not backward:
            continue
        t1, t1_column = forward
        t2_id = (t1_column.get('data') or {}).get('other_table_id')
        t2 = table_by_id.get(t2_id)
        if not t2:
            continue
        t2_column = next((column for table, column in columns if table['id'] == t2_id and column != t1_column), None)
        if not t2_column:
            continue
        definition = [
            link_id, t1['id'], t2['id'],
            t1_column['key'], t1_column['name'], (t1_column.get('data') or {}).get('display_column_key'),
            t2_column['key'], t2_column['name'], (t2_column.get('data') or {}).get('display_column_key'),
        ]
        wrote_relation = False
        for row in table_rows[t1['id']]:
            for linked in row.get(t1_column['name']) or []:
                sheet.append(definition + [row['_id'], linked.get('row_id')])
                wrote_relation = True
        if not wrote_relation:
            sheet.append(definition + [None, None])


def _write_key_values(sheet, values):
    _write_header(sheet, ['key', 'part', 'value'])
    for key, value in values.items():
        encoded = json.dumps(value, ensure_ascii=False, separators=(',', ':'))
        for part, chunk in enumerate(_chunks(encoded), 1):
            sheet.append([_string_cell(sheet, key), part, _string_cell(sheet, chunk)])


def _read_key_values(sheet):
    chunks = {}
    for key_cell, part_cell, value_cell, *_ in sheet.iter_rows(min_row=2):
        key = key_cell.value
        if not key:
            continue
        if value_cell.data_type == 'f':
            raise MetadataBackupError(f'Formula not allowed for {key}')
        chunks.setdefault(key, {})[part_cell.value] = value_cell.value
    values = {}
    for key, parts in chunks.items():
        try:
            values[key] = json.loads(_join_chunks(parts))
        except (TypeError, ValueError) as error:
            raise MetadataBackupError(f'Invalid value for {key}') from error
    return values


def _read_schema(sheet):
    raw_columns = {}
    allowed_modes = {'identity', 'preserve', 'conditional', 'restore'}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in values):
            continue
        if len(values) < len(SCHEMA_HEADERS):
            raise MetadataBackupError('Invalid schema row')
        table_id, table_name, column_index, key, name, column_type, mode, data_part, data = values[:9]
        if not all((table_id, table_name, column_index, key, name, column_type)) or mode not in allowed_modes:
            raise MetadataBackupError('Invalid schema row')
        identity = (table_id, key)
        metadata = (table_name, int(column_index), name, column_type, mode)
        item = raw_columns.setdefault(identity, {'metadata': metadata, 'parts': {}})
        if item['metadata'] != metadata or data_part in item['parts']:
            raise MetadataBackupError(f'Invalid schema for column {key}')
        item['parts'][data_part] = data
    schema = {}
    for (table_id, key), item in raw_columns.items():
        table_name, column_index, name, column_type, mode = item['metadata']
        try:
            column_data = json.loads(_join_chunks(item['parts']))
        except (TypeError, ValueError) as error:
            raise MetadataBackupError(f'Invalid schema for column {key}') from error
        schema.setdefault(table_id, []).append({
            'table_id': table_id, 'table_name': table_name, 'column_index': column_index,
            'key': key, 'name': name, 'type': column_type, 'data': column_data, 'mode': mode,
        })
    for columns in schema.values():
        columns.sort(key=lambda column: column['column_index'])
        keys = [column['key'] for column in columns]
        indexes = [column['column_index'] for column in columns]
        if len(keys) != len(set(keys)) or indexes != list(range(1, len(columns) + 1)):
            raise MetadataBackupError('Invalid schema columns')
    return schema


def _validate_schema(schema):
    main_columns = schema[MAIN_TABLE_ID]
    if main_columns[0]['table_name'] != MAIN_TABLE_NAME:
        raise MetadataBackupError('Invalid metadata table')
    tag_tables = [columns for table_id, columns in schema.items() if table_id != MAIN_TABLE_ID]
    if len(tag_tables) > 1:
        raise MetadataBackupError('Invalid tags schema')
    for table_id, columns in schema.items():
        table = {'id': table_id, 'name': columns[0]['table_name']}
        if table_id != MAIN_TABLE_ID and table['name'] != TAGS_TABLE_NAME:
            raise MetadataBackupError('Unsupported metadata table')
        names = [column['name'] for column in columns]
        if len(names) != len(set(names)):
            raise MetadataBackupError('Duplicate column name')
        for column in columns:
            if column['type'] not in SUPPORTED_COLUMN_TYPES:
                raise MetadataBackupError(f'Unsupported column type {column["type"]}')
            if column['mode'] != _column_mode(table, column):
                raise MetadataBackupError(f'Invalid restore mode for column {column["key"]}')


def _read_data_sheet(sheet, columns, long_values):
    headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    if list(headers[:len(columns)]) != [column['name'] for column in columns]:
        raise MetadataBackupError(f'Invalid headers in {sheet.title}')
    rows = []
    for cells in sheet.iter_rows(min_row=2, max_col=len(columns)):
        if not any(cell.value is not None for cell in cells):
            continue
        row = {}
        for column, cell in zip(columns, cells):
            if cell.data_type == 'f':
                raise MetadataBackupError(f'Formulas are not allowed in {sheet.title}')
            row[column['key']] = _decode_value(column, cell.value)
        if not row.get('_id'):
            raise MetadataBackupError(f'Row id missing in {sheet.title}')
        for column in columns:
            long_value = long_values.get((column['table_id'], row['_id'], column['key']))
            if long_value is not None:
                row[column['key']] = _decode_value(column, long_value)
        rows.append(row)
    return rows


def _decode_value(column, value):
    if value is None:
        return None
    column_type = column['type']
    if column_type in JSON_TYPES:
        try:
            return json.loads(value)
        except (TypeError, ValueError) as error:
            raise MetadataBackupError(f'Invalid JSON value for {column["name"]}') from error
    if column_type in {'number', 'rate', 'duration'}:
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise MetadataBackupError(f'Invalid number for {column["name"]}')
        return value
    if column_type == 'checkbox':
        if not isinstance(value, bool):
            raise MetadataBackupError(f'Invalid checkbox for {column["name"]}')
        return value
    if not isinstance(value, str):
        raise MetadataBackupError(f'Invalid text for {column["name"]}')
    return value


def _validate_unique_row_ids(rows, table_name):
    row_ids = [row['_id'] for row in rows]
    if len(row_ids) != len(set(row_ids)):
        raise MetadataBackupError(f'Duplicate row id in {table_name}')


def _api_column(column):
    return {
        'key': column['key'], 'name': column['name'],
        'type': column['type'], 'data': column.get('data'),
    }


def _read_links(sheet):
    definitions = {}
    relations = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in values):
            continue
        if len(values) < len(LINK_DEFINITION_HEADERS):
            raise MetadataBackupError('Invalid link row')
        link_id = values[0]
        definition = tuple(values[:len(LINK_DEFINITION_HEADERS)])
        if not link_id or (link_id in definitions and definitions[link_id] != definition):
            raise MetadataBackupError('Invalid link definition')
        definitions[link_id] = definition
        relation = tuple(values[len(LINK_DEFINITION_HEADERS):len(LINK_HEADERS)])
        relation += (None,) * (len(LINK_RELATION_HEADERS) - len(relation))
        row_id, other_row_id = relation
        if bool(row_id) != bool(other_row_id):
            raise MetadataBackupError('Invalid link relation')
        if row_id:
            relations.setdefault(link_id, {}).setdefault(row_id, []).append(other_row_id)

    links = []
    for link_id, definition in definitions.items():
        _, t1_id, t2_id, t1_key, t1_name, t1_display, t2_key, t2_name, t2_display = definition
        links.append({
            'link_id': link_id, 't1_id': t1_id, 't2_id': t2_id,
            't1_column': {'key': t1_key, 'name': t1_name, 'display_column_key': t1_display},
            't2_column': {'key': t2_key, 'name': t2_name, 'display_column_key': t2_display},
            'row_id_map': relations.get(link_id, {}),
        })
    return links


def _read_long_values(sheet):
    chunks = {}
    for cells in sheet.iter_rows(min_row=2, max_col=len(VALUE_HEADERS)):
        if not any(cell.value is not None for cell in cells):
            continue
        table_id, row_id, column_key, part = [cell.value for cell in cells[:4]]
        value_cell = cells[4]
        if not table_id or not row_id or not column_key or not part or value_cell.data_type == 'f':
            raise MetadataBackupError('Invalid long metadata value')
        identity = (table_id, row_id, column_key)
        if part in chunks.setdefault(identity, {}):
            raise MetadataBackupError('Duplicate long metadata value part')
        chunks[identity][part] = value_cell.value
    return {identity: _join_chunks(parts) for identity, parts in chunks.items()}


def _chunks(value):
    return [value[index:index + CELL_CHUNK_SIZE] for index in range(0, len(value), CELL_CHUNK_SIZE)] or ['']


def _join_chunks(parts):
    indexes = sorted(parts)
    if indexes != list(range(1, len(parts) + 1)):
        raise MetadataBackupError('Metadata value parts are incomplete')
    return ''.join(parts[index] or '' for index in indexes)
